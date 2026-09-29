#!/usr/bin/env python3
import sys
import json
import re
import openpyxl

MONTH_NAME_MAP = {
    'jan': '01', 'janeiro': '01', 'january': '01',
    'fev': '02', 'fevereiro': '02', 'february': '02',
    'mar': '03', 'marco': '03', 'março': '03', 'march': '03',
    'abr': '04', 'abril': '04', 'april': '04',
    'mai': '05', 'maio': '05', 'may': '05',
    'jun': '06', 'junho': '06', 'june': '06',
    'jul': '07', 'julho': '07', 'july': '07',
    'ago': '08', 'agosto': '08', 'august': '08',
    'set': '09', 'setembro': '09', 'sept': '09', 'september': '09', 'setout': '09',
    'out': '10', 'outubro': '10', 'october': '10',
    'nov': '11', 'novembro': '11', 'november': '11',
    'dez': '12', 'dezembro': '12', 'december': '12',
}

def clean_num(val):
    if val is None:
        return 0
    if isinstance(val, (int, float)):
        return round(float(val), 2)
    s = str(val).strip().replace(' ', '').replace(',', '.')
    s = re.sub(r'[^0-9.]', '', s)
    try:
        return round(float(s), 2) if s else 0
    except:
        return 0

def parse_month_year(sheet_name, row2_val, sheet_idx):
    # Year estimation starting from 2021
    # Sheet index spans from 2021 to 2026
    year = 2021
    if sheet_idx >= 45:
        year = 2026
    elif sheet_idx >= 33:
        year = 2025
    elif sheet_idx >= 21:
        year = 2024
    elif sheet_idx >= 9:
        year = 2023
    elif sheet_idx >= 4:
        year = 2022

    # Check for explicit year in sheet name or row2
    text = f"{sheet_name} {str(row2_val or '')}"
    match_year = re.search(r'202[1-6]', text)
    if match_year:
        year = int(match_year.group(0))

    # Check for month name
    month_num = '01'
    text_lower = text.lower()
    for m_name, m_val in MONTH_NAME_MAP.items():
        if m_name in text_lower:
            month_num = m_val
            break

    return f"{year}-{month_num}"

def parse_workbook(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    months_data = []

    for idx, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        row2_val = ws.cell(2, 1).value
        month_key = parse_month_year(sheet_name, row2_val, idx)

        # Parse Row 6 Income & Rules
        gross_salary = clean_num(ws.cell(6, 1).value)
        inss = clean_num(ws.cell(6, 2).value)
        tithe = clean_num(ws.cell(6, 3).value)
        net_income = clean_num(ws.cell(6, 4).value)

        current_section = None
        sections = {
            'Dividas': [],
            'Casa': [],
            'Familia': [],
            'Outros': [],
        }

        for r in range(8, ws.max_row + 1):
            c1 = ws.cell(r, 1).value
            c2 = ws.cell(r, 2).value
            c3 = ws.cell(r, 3).value
            c4 = ws.cell(r, 4).value

            if c1 is None:
                continue

            c1_str = str(c1).strip()
            if not c1_str:
                continue

            # Check section header
            if c1_str.lower().startswith('divida'):
                current_section = 'Dividas'
                continue
            elif c1_str.lower().startswith('casa'):
                current_section = 'Casa'
                continue
            elif c1_str.lower().startswith('famili'):
                current_section = 'Familia'
                continue
            elif c1_str.lower() in ['total', 'subtotal', 'dispesas', 'saldo']:
                continue

            val = clean_num(c2)
            if val > 0 and current_section:
                obs = str(c3 or c4 or '').strip()
                sections[current_section].append({
                    'item': c1_str,
                    'amount': val,
                    'obs': obs,
                })

        months_data.append({
            'sheet_name': sheet_name,
            'month': month_key,
            'gross_salary': gross_salary,
            'inss': inss,
            'tithe': tithe,
            'net_income': net_income,
            'sections': sections,
        })

    return months_data

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: parse_excel.py <path_to_xlsx>")
        sys.exit(1)

    data = parse_workbook(sys.argv[1])
    print(json.dumps(data, ensure_ascii=False, indent=2))
